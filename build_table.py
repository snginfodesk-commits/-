# -*- coding: utf-8 -*-
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "השוואת פרויקטים"
ws.sheet_view.rightToLeft = True

NUM_PROJECTS = 5  # עמודות ריקות למילוי

# צבעים
CAT_FILL = PatternFill("solid", fgColor="1F4E78")   # כותרת קטגוריה
HEAD_FILL = PatternFill("solid", fgColor="2E75B6")  # שורת כותרת פרויקטים
FIELD_FILL = PatternFill("solid", fgColor="D9E1F2") # עמודת שמות שדות
ALT_FILL = PatternFill("solid", fgColor="F2F6FC")   # פסים מתחלפים

WHITE = Font(name="Arial", size=11, bold=True, color="FFFFFF")
CAT_FONT = Font(name="Arial", size=12, bold=True, color="FFFFFF")
FIELD_FONT = Font(name="Arial", size=11, bold=True, color="1F2D3D")

thin = Side(style="thin", color="BFBFBF")
border = Border(left=thin, right=thin, top=thin, bottom=thin)

center = Alignment(horizontal="center", vertical="center", wrap_text=True)
right = Alignment(horizontal="right", vertical="center", wrap_text=True)

# מבנה: רשימת (קטגוריה, [שדות])
structure = [
    ("סטטוס הפרויקט", [
        "תאריך בדיקה",
        "שם היזם",
        "שם הפרויקט (כתובת)",
        "מחיר הדירה",
        "סטטוס היתר בנייה (יש היתר / החלטת ועדה / בבקשה)",
        "מועד מסירה משוער (מתי מוכן)",
    ]),
    ("נתונים פיזיים ותכנוניים", [
        "מספר בניינים בפרויקט",
        "סך הכל דירות בפרויקט",
        'דירת 3 חדרים: שטח בנוי (מ"ר)',
        'דירת 3 חדרים: שטח מרפסת (מ"ר)',
        "קומה וכיווני אוויר",
        "הצמדות (חניה / מחסן)",
    ]),
    ("נתונים כספיים ותנאים", [
        "תנאי תשלום (למשל: 20/80)",
        "בנק מלווה וערבות חוק מכר",
        "מנגנון הצמדה למדד (פטור / חסום / חלקי)",
        'פטורים והנחות (פטור משכ"ט עו"ד / שדרוגי מפרט)',
    ]),
    ("ניתוח שוק ואזור", [
        'מחיר ממוצע למ"ר באזור',
        "צפי שכירות חודשית",
        "מרחק מתחבורה ציבורית (רכבת קלה / מטרו / תחנה מרכזית)",
    ]),
]

total_cols = 1 + NUM_PROJECTS
last_col_letter = get_column_letter(total_cols)

row = 1
# שורת כותרת ראשית - פרויקטים
ws.cell(row=row, column=1, value="פרמטר / פרויקט")
for c in range(2, total_cols + 1):
    ws.cell(row=row, column=c, value=f"פרויקט {c-1}")
for c in range(1, total_cols + 1):
    cell = ws.cell(row=row, column=c)
    cell.fill = HEAD_FILL
    cell.font = WHITE
    cell.alignment = center
    cell.border = border
ws.row_dimensions[row].height = 28
row += 1

for cat, fields in structure:
    # שורת קטגוריה (ממוזגת לכל הרוחב)
    ws.cell(row=row, column=1, value=cat)
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=total_cols)
    cell = ws.cell(row=row, column=1)
    cell.fill = CAT_FILL
    cell.font = CAT_FONT
    cell.alignment = right
    cell.border = border
    ws.row_dimensions[row].height = 26
    row += 1

    for i, field in enumerate(fields):
        ws.cell(row=row, column=1, value=field)
        fc = ws.cell(row=row, column=1)
        fc.fill = FIELD_FILL
        fc.font = FIELD_FONT
        fc.alignment = right
        fc.border = border
        for c in range(2, total_cols + 1):
            vc = ws.cell(row=row, column=c, value="")
            vc.fill = ALT_FILL if i % 2 == 0 else PatternFill("solid", fgColor="FFFFFF")
            vc.alignment = right
            vc.border = border
        ws.row_dimensions[row].height = 34
        row += 1

# רוחב עמודות
ws.column_dimensions["A"].width = 42
for c in range(2, total_cols + 1):
    ws.column_dimensions[get_column_letter(c)].width = 22

# הקפאת חלוניות - עמודת השדות + שורת הכותרת
ws.freeze_panes = "B2"

wb.save("טבלת_מעקב_פרויקטים.xlsx")
print("saved")
